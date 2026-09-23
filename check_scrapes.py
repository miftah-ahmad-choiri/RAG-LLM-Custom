import os
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# Ensure we use identical slugify logic from excel_parser.py
def slugify(text: str, max_len: int = 60) -> str:
    text = text.strip().lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    text = text.strip("-")
    return text[:max_len].rstrip("-")

def _top_index(index: str) -> str:
    return index.split(".")[0]

def main():
    import sys
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")

    workspace_dir = Path(".")
    excel_in_path = workspace_dir / "ceph-scraper" / "output" / "excel" / "IBM-Ceph-TOC-20260922_231027.xlsx"
    excel_out_path = workspace_dir / "ceph-scraper" / "output" / "excel" / "IBM-Ceph-TOC-20260922_231027_checked.xlsx"
    md_root = workspace_dir / "ceph-scraper" / "output" / "md" / "20260922_234455-IBM-Ceph-TOC-20260922_231027"

    print(f"Loading workbook: {excel_in_path}")
    wb = load_workbook(str(excel_in_path))
    ws = wb["TOC"]

    # First pass: Parse and build mapping of all indices to descriptions and urls
    entries = []
    section_map = {} # ti -> section_description

    # Row 1 is header: Index, Description, URL
    for row_idx in range(2, ws.max_row + 1):
        idx = ws.cell(row=row_idx, column=1).value
        desc = ws.cell(row=row_idx, column=2).value
        url_formula = ws.cell(row=row_idx, column=3).value

        if idx is None:
            continue

        idx = str(idx).strip()
        desc = str(desc).strip() if desc else idx
        ti = _top_index(idx)

        # If it is a top-level section page (no dots in index)
        if "." not in idx:
            section_map[idx] = desc

        entries.append({
            "row_idx": row_idx,
            "index": idx,
            "description": desc,
            "top_index": ti,
        })

    print(f"Parsed {len(entries)} entries from Excel. Found {len(section_map)} top-level sections.")

    # Boilerplate signatures to check
    boilerplate_patterns = [
        r"was this topic helpful\?",
        r"focus sentinel",
        r"rate this content",
        r"thank you for your feedback",
        r"together, we can continue to improve ibm documentation",
        r"return to topic",
        r"thank you for your submission",
        r"submissions are limited to 1 per day per topic",
        r"error submitting rating",
        r"there has been an error sending your feedback to the team",
        r"please try again later",
        r"change version",
        r"active loading indicator",
    ]

    # Specific boilerplate lines to strip for content-emptiness checks
    boilerplate_exact_lines = {
        "change version", "9.9.1", "9.9.0", "8.1.0", "8.0.0", "7.1.0", "7.0.0", "6.1.0", "5.3.0",
        "was this topic helpful?", "focus sentinel", "close", "rate this content",
        "thank you for your feedback!", "together, we can continue to improve ibm documentation.",
        "return to topic", "thank you for your submission.", "submissions are limited to 1 per day per topic.",
        "error submitting rating", "there has been an error sending your feedback to the team. your comment was saved locally, if not in an incognito browser, and will be available when attempting to submit feedback again.",
        "please try again later.", "loading", "active loading indicator"
    }

    # Style definitions for column D
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    fill_header = PatternFill(start_color="0F62FE", end_color="0F62FE", fill_type="solid")
    
    # Fills for different status notes
    fill_failed = PatternFill(start_color="FFD0D0", end_color="FFD0D0", fill_type="solid") # light red
    fill_incomplete = PatternFill(start_color="FFEAA7", end_color="FFEAA7", fill_type="solid") # light yellow
    fill_clean = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid") # light green
    fill_missing = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid") # light gray

    # Set Header for Column D
    header_cell = ws.cell(row=1, column=4)
    header_cell.value = "Scrape Status & Note"
    header_cell.font = font_header
    header_cell.fill = fill_header
    header_cell.alignment = Alignment(horizontal="center", vertical="center")

    stats = {
        "missing": 0,
        "empty": 0,
        "failed_boilerplate_only": 0,
        "incomplete_with_boilerplate": 0,
        "clean": 0,
    }

    failed_details = []
    suspicious_details = []

    print("Checking markdown files on disk and applying notes...")

    for entry in entries:
        idx = entry["index"]
        desc = entry["description"]
        ti = entry["top_index"]
        row_idx = entry["row_idx"]

        # Determine section folder name
        sec_desc = section_map.get(ti)
        if not sec_desc:
            # Fallback
            sec_folder_name = ti
        else:
            sec_folder_name = f"{ti}-{slugify(sec_desc)}"

        # Determine file path
        file_slug = slugify(desc)
        filename = f"{idx}-{file_slug}.md"
        filepath = md_root / sec_folder_name / filename

        note = ""
        cell_fill = None

        if not filepath.exists():
            note = f"⚠️ File not found on disk: {sec_folder_name}/{filename}"
            cell_fill = fill_missing
            stats["missing"] += 1
        else:
            # Read file and check content
            try:
                content = filepath.read_text(encoding="utf-8")
            except Exception as e:
                content = ""
                note = f"❌ Error reading file: {str(e)}"
                cell_fill = fill_missing
                stats["missing"] += 1

            if not note:
                lines = [line.strip() for line in content.split("\n")]
                
                # Check for boilerplate presence
                has_boilerplate = False
                found_boilerplates = []
                content_lower = content.lower()
                for pat in boilerplate_patterns:
                    if re.search(pat, content_lower):
                        has_boilerplate = True
                        found_boilerplates.append(pat.replace("\\?", "?"))

                # Filter lines to evaluate real content
                real_lines = []
                for line in lines:
                    # Clean/normalize the line to check against boilerplate phrases
                    cleaned_line = line.lstrip("#").strip()
                    cleaned_lower = cleaned_line.lower()
                    
                    # Skip metadata header/footer elements
                    if line.startswith("# ") or line.startswith("> **Index:**") or line.startswith("> **Source:**") or line == "---" or line.startswith("*Source:"):
                        continue
                        
                    # Skip standard blank lines
                    if not line:
                        continue
                        
                    # Skip exact boilerplate matches
                    if cleaned_lower in boilerplate_exact_lines:
                        continue
                        
                    # Substring matches for boilerplate and system warnings
                    boilerplate_substrings = [
                        "focus sentinel",
                        "was this topic helpful",
                        "rate this content",
                        "thank you for your feedback",
                        "together, we can continue to improve ibm documentation",
                        "return to topic",
                        "thank you for your submission",
                        "submissions are limited to 1 per day per topic",
                        "error submitting rating",
                        "there has been an error sending your feedback",
                        "please try again later",
                        "change version",
                        "active loading indicator",
                        "loading",
                        "close",
                        "get hands-on experience with ibm tech",
                        "?lit$",
                        "docs offline",
                        "security update required for ibm docs offline",
                        "continue download",
                        "visit the docs offline page",
                        "view the security bulletin"
                    ]
                    
                    if any(sub in cleaned_lower for sub in boilerplate_substrings):
                        continue

                    # If it passed all filters, it is a line of real content
                    real_lines.append(line)

                # Classify the file based on real content lines
                real_content_len = len(" ".join(real_lines).strip())
                
                if len(content.strip()) == 0:
                    note = "⚠️ Empty file (0 bytes)"
                    cell_fill = fill_missing
                    stats["empty"] += 1
                elif real_content_len < 10:
                    # Almost no real content remains, meaning it's a failed scrape with just boilerplate!
                    note = "🚨 Failed Scrape (Boilerplate footer/headers only, no actual content)"
                    cell_fill = fill_failed
                    stats["failed_boilerplate_only"] += 1
                    failed_details.append({
                        "index": idx,
                        "description": desc,
                        "file": f"{sec_folder_name}/{filename}",
                        "len": real_content_len,
                    })
                elif has_boilerplate:
                    # Has boilerplate, but also has substantial real content
                    note = f"⚠️ Incomplete Scrape (Contains boilerplate footer, but has {len(real_lines)} lines of real content)"
                    cell_fill = fill_incomplete
                    stats["incomplete_with_boilerplate"] += 1
                    if real_content_len < 200:
                        suspicious_details.append({
                            "index": idx,
                            "description": desc,
                            "file": f"{sec_folder_name}/{filename}",
                            "len": real_content_len,
                        })
                else:
                    note = f"✅ Clean Scrape (No boilerplate found, has {len(real_lines)} lines of real content)"
                    cell_fill = fill_clean
                    stats["clean"] += 1

        # Write to column D
        cell = ws.cell(row=row_idx, column=4)
        cell.value = note
        cell.font = font_data
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if cell_fill:
            cell.fill = cell_fill

    # Set column width for column D
    ws.column_dimensions["D"].width = 50

    print("Saving modified workbook...")
    wb.save(str(excel_out_path))
    wb.close()

    print("\n" + "="*50)
    print("SCRAPE VALIDATION SUMMARY:")
    print("="*50)
    print(f"Clean Scrapes (No Boilerplate):            {stats['clean']}")
    print(f"Incomplete Scrapes (Contains boilerplate): {stats['incomplete_with_boilerplate']}")
    print(f"Failed Scrapes (Boilerplate only):         {stats['failed_boilerplate_only']}")
    print(f"Empty Files:                               {stats['empty']}")
    print(f"Missing Files:                             {stats['missing']}")
    print(f"Total checked:                             {len(entries)}")
    print("="*50)
    
    if failed_details:
        print("\n🚨 DETAILED FAILED SCRAPES (Boilerplate only):")
        print("-" * 50)
        for d in failed_details:
            print(f"Index {d['index']}: {d['description']}")
            print(f"  File: {d['file']} (Real content length: {d['len']} chars)")
            print("-" * 50)
            
    if suspicious_details:
        print("\n⚠️ DETAILED SUSPICIOUS/SHORT INCOMPLETE SCRAPES (< 200 chars real content):")
        print("-" * 50)
        # Sort by length ascending
        suspicious_details.sort(key=lambda x: x["len"])
        for d in suspicious_details[:20]: # show top 20 shortest
            print(f"Index {d['index']}: {d['description']}")
            print(f"  File: {d['file']} (Real content length: {d['len']} chars)")
            print("-" * 50)

    print(f"Output saved to: {excel_out_path}")

if __name__ == "__main__":
    main()
